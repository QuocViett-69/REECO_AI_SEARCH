from pathlib import Path
import json
import re
import zipfile
from PIL import Image, ImageChops
from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[1]
SOURCE = ROOT / 'docs/bao-cao-chung-gui-quan-ly-theo-role.md'
OUTPUT = ROOT / 'docs/bao-cao-reeco-ai-search.docx'
text = SOURCE.read_text(encoding='utf-8')
figures = json.loads((BASE / 'figures.json').read_text(encoding='utf-8'))

def clean(value):
    value = re.sub(r'\[([^]]+)\]\([^)]+\)',r'\1',value)
    return value.replace('**','').replace('`','').strip()

def paragraphs(value):
    return [clean(p) for p in re.split(r'\n\s*\n',value.strip())
            if p.strip() and not p.lstrip().startswith(('**Sơ đồ','###','##'))]

def between(start,end):
    return text.split(start,1)[1].split(end,1)[0]

def markdown_table(value):
    rows=[]
    for line in value.splitlines():
        if line.startswith('|') and not re.fullmatch(r'[| :\-]+',line):
            rows.append([clean(v) for v in line.strip('|').split('|')])
    return rows

# Trim the paper variants only; the editable draw.io originals are unchanged.
for item in figures:
    path=BASE/'assets'/f"figure-{item['number']:02d}.png"
    image=Image.open(path).convert('RGB')
    diff=ImageChops.difference(image,Image.new('RGB',image.size,'white'))
    mask=diff.convert('L').point(lambda p:255 if p>12 else 0)
    bbox=mask.getbbox()
    assert bbox, path
    pad=28
    box=(max(0,bbox[0]-pad),max(0,bbox[1]-pad),min(image.width,bbox[2]+pad),min(image.height,bbox[3]+pad))
    image.crop(box).save(path,dpi=(240,240))

doc=Document()
doc.core_properties.title='Báo cáo hệ thống REECO AI Search'
doc.core_properties.subject='Phân công và báo cáo theo FE BE BA AI Engineer'
doc.core_properties.author='Nhóm đồ án REECO AI Search'
doc.core_properties.keywords='REECO, AI Search, báo cáo, FE, BE, BA, AI Engineer'
doc.core_properties.comments=''

styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','List Bullet','TOC 1']:
    if name not in styles: continue
    style=styles[name]
    style.font.name='Arial'
    style.font.color.rgb=RGBColor(0,0,0)
    rp=style.element.get_or_add_rPr()
    fonts=rp.find(qn('w:rFonts'))
    if fonts is None: fonts=OxmlElement('w:rFonts');rp.append(fonts)
    for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Arial')
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        fonts.attrib.pop(qn('w:'+attr),None)
    color=rp.find(qn('w:color'))
    if color is not None:
        color.set(qn('w:val'),'000000')
        for attr in ['themeColor','themeTint','themeShade']:color.attrib.pop(qn('w:'+attr),None)
normal=styles['Normal']
normal.font.size=Pt(11.25)
normal.paragraph_format.line_spacing=1.16
normal.paragraph_format.space_after=Pt(8)
normal.paragraph_format.widow_control=True
for name,size,before,after in [('Title',28,0,18),('Subtitle',14,0,10),('Heading 1',19,0,14),('Heading 2',13.5,12,8),('Heading 3',11.5,10,6)]:
    style=styles[name];style.font.size=Pt(size);style.font.bold=name!='Subtitle'
    pf=style.paragraph_format;pf.space_before=Pt(before);pf.space_after=Pt(after);pf.keep_with_next=True
    pp=style.element.find(qn('w:pPr'))
    if pp is not None:
        border=pp.find(qn('w:pBdr'))
        if border is not None:pp.remove(border)
styles['Caption'].font.size=Pt(10)
styles['Caption'].font.italic=False
styles['Subtitle'].font.italic=False
styles['Caption'].paragraph_format.space_before=Pt(8)
styles['Caption'].paragraph_format.space_after=Pt(8)

def field(p,instruction,cached=''):
    r=p.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r._r.append(begin)
    instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' '+instruction+' ';p.add_run()._r.append(instr)
    sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');p.add_run()._r.append(sep)
    p.add_run(cached)
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');p.add_run()._r.append(end)

def configure(section,landscape=False):
    section.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width=Cm(29.7 if landscape else 21)
    section.page_height=Cm(21 if landscape else 29.7)
    section.left_margin=section.right_margin=Cm(1.35 if landscape else 2)
    section.top_margin=Cm(1.2 if landscape else 1.8)
    section.bottom_margin=Cm(1.1 if landscape else 1.7)
    section.header_distance=Cm(.6);section.footer_distance=Cm(.6)
    # Page footer is centered independent of page orientation.
    p=section.footer.paragraphs[0]
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(0)
    if not p.text and not p._p.xpath('.//w:fldChar'):
        r=p.add_run('Trang ');r.font.size=Pt(9)
        field(p,'PAGE','1')
    for r in p.runs:r.font.size=Pt(9)

configure(doc.sections[0])
doc.sections[0].different_first_page_header_footer=True
bookmark_id=1
def bookmark(p,name):
    global bookmark_id
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bookmark_id));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark_id))
    p._p.insert(0,start);p._p.append(end);bookmark_id+=1

def para(value,bold=False,italic=False,size=None):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    r=p.add_run(clean(value));r.bold=bold;r.italic=italic
    if size:r.font.size=Pt(size)
    return p

def heading(value,level=1,newpage=False):
    p=doc.add_paragraph(value,style=f'Heading {level}')
    if newpage:p.paragraph_format.page_break_before=True
    return p

def newsection(landscape):
    section=doc.add_section(WD_SECTION_START.NEW_PAGE)
    configure(section,landscape)
    section.different_first_page_header_footer=False
    return section

def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+name)
        for key,val in [('val','single'),('sz','5'),('color','D9D9D9')]:el.set(qn('w:'+key),val)
        borders.append(el)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for name in ['top','left','bottom','right']:
        el=OxmlElement('w:'+name);el.set(qn('w:w'),'110');el.set(qn('w:type'),'dxa');margins.append(el)
    pr.append(margins)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    allrows=[headers]+rows
    for index,data in enumerate(allrows):
        row=t.rows[0] if index==0 else t.add_row()
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for ci,(cell,value,w) in enumerate(zip(row.cells,data,widths)):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'DCE6F1' if index==0 else ('F4F7FA' if index%2==0 else 'FFFFFF'));tcpr.append(shd)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.08
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ci==0 and len(headers)>2 else WD_ALIGN_PARAGRAPH.LEFT
            r=p.add_run(clean(value));r.font.size=Pt(10);r.bold=index==0;r.font.color.rgb=RGBColor(0,0,0)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(4);p.add_run().font.size=Pt(4)
    return t

def addfigure(number,section_number,subnumber):
    item=figures[number-1]
    newsection(True)
    p=doc.add_paragraph(f'{section_number} {subnumber} {item["caption"]}',style='Heading 2')
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(8)
    image=Image.open(BASE/'assets'/f'figure-{number:02d}.png')
    # Keep all figures inside one landscape A4 page without changing aspect ratio.
    max_w,max_h=27.0,16.0
    ratio=min(max_w/image.width,max_h/image.height)
    w,h=image.width*ratio,image.height*ratio
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
    p.paragraph_format.keep_with_next=True
    drawing=p.add_run().add_picture(str(BASE/'assets'/f'figure-{number:02d}.png'),width=Cm(w),height=Cm(h))
    drawing._inline.docPr.set('descr',item['caption'])
    caption=doc.add_paragraph(f'Hình {number}. {item["caption"]}',style='Caption')
    caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
    bookmark(caption,f'figure{number}')

# Cover page.
p=doc.add_paragraph('REECO AI Search',style='Subtitle');p.paragraph_format.space_before=Cm(2.4)
doc.add_paragraph('Báo cáo hệ thống\nREECO AI Search',style='Title')
doc.add_paragraph('Phân công và triển khai theo bốn vai trò',style='Subtitle')
p=para('Frontend Developer và Backend Developer\nBusiness Analyst và AI Engineer kiêm xử lý dữ liệu',size=12)
p.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.add_paragraph().paragraph_format.space_after=Pt(22)
for label,value in [('Đơn vị thực hiện','Nhóm đồ án REECO AI Search'),('Người nhận','Quản lý dự án'),('Ngày lập','26 tháng 10 năm 2026'),('Thứ tự trình bày','FE → BE → BA → AI Engineer')]:
    p=doc.add_paragraph();p.add_run(label+'  ').bold=True;p.add_run(value)
para('Báo cáo trình bày phạm vi hệ thống, trách nhiệm từng vai trò, thiết kế bằng sơ đồ và kết quả đánh giá đã lưu. Mục tiêu là thống nhất nội dung trước khi nghiệm thu và trình diễn.',size=11.25)

# Navigation page with real editable TOC and page references.
doc.add_page_break()
p=doc.add_paragraph('Mục lục');p.style='Subtitle';p.runs[0].bold=True
field(doc.add_paragraph(),'TOC \\o "1-1" \\h \\z \\u','')
p=doc.add_paragraph('Danh mục hình');p.style='Subtitle';p.paragraph_format.space_before=Pt(18);p.runs[0].bold=True
for item in figures:
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.5),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(f'Hình {item["number"]}. {item["caption"]}\t').font.size=Pt(10)
    field(p,f'PAGEREF figure{item["number"]} \\h','')

heading('1 Tổng quan và phân công',newpage=True)
intro=between('Kính gửi Quản lý dự án,','### 1. Frontend Developer')
for p in paragraphs(intro):para(p)
heading('1 1 Trách nhiệm và sản phẩm bàn giao',2)
rows=[
['FE','Client và admin; routing; state/session phía trình duyệt; tích hợp API và hiển thị kết quả AI.','Sơ đồ màn hình, giao diện, kiểm thử frontend và demo.'],
['BE','API, database, xác thực và phân quyền; giá, đơn, VNPay, thông báo và vận hành.','Kiến trúc, ERD, API contract, sequence và kiểm thử backend.'],
['BA','Bài toán, yêu cầu, use case, quy tắc nghiệp vụ và tiêu chí chấp nhận.','Đặc tả, ma trận truy vết, kế hoạch UAT và báo cáo chung.'],
['AI và Data','Chuẩn hóa dữ liệu, search document, embedding, retrieval, kiểm chứng, reranker và explanation.','Pipeline dữ liệu và tìm kiếm, benchmark, bằng chứng và giới hạn.'],
]
table(['Vai trò','Phạm vi phụ trách','Sản phẩm bàn giao'],rows,[2.2,8.2,6.6])
para('FE phụ trách cả ứng dụng khách hàng và quản trị. AI Engineer phụ trách thuật toán tìm kiếm kể cả service viết bằng NestJS; BE phụ trách hạ tầng tích hợp và database dùng chung. BA tổng hợp báo cáo; mỗi vai trò tự viết và xác nhận phần chuyên môn của mình.')

heading('1 2 Phân công viết báo cáo đồ án',2,newpage=True)
source_rows=markdown_table(between('### Phân công viết báo cáo đồ án','### Ranh giới phối hợp'))
table(source_rows[0],source_rows[1:],[5.0,12.0])
heading('1 3 Phối hợp giữa các vai trò',2)
for line in between('### Ranh giới phối hợp','## B.').splitlines():
    if line.startswith('- '):
        p=doc.add_paragraph(clean(line[2:]),style='List Bullet');p.paragraph_format.space_after=Pt(7)

roles=[
('2 Frontend Developer','### 1. Frontend Developer','### 2. Backend Developer',[1]),
('3 Backend Developer','### 2. Backend Developer','### 3. Business Analyst',[2,3,4,5,6,7]),
('4 Business Analyst','### 3. Business Analyst','### 4. AI Engineer',[8,9]),
('5 AI Engineer và xử lý dữ liệu','### 4. AI Engineer','### Kết luận và bước tiếp theo',[10,11,12]),
]
for index,(title,start,end,numbers) in enumerate(roles):
    if index==0:heading(title,newpage=True)
    else:newsection(False);heading(title)
    ps=paragraphs(between(start,end).split('\n',1)[1])
    for pi,p in enumerate(ps):
        if index==3 and pi==3:continue
        para(p)
    if index==2:
        heading('4 1 Tiêu chí chấp nhận cần đối chiếu',2)
        table(['Nhóm','Tiêu chí'],[
          ['Tìm kiếm','Điều kiện kỹ thuật bắt buộc đang áp dụng phải có bằng chứng; thiếu thông tin cần thể hiện chưa xác nhận.'],
          ['Đặt hàng','Giá và tổng tiền do server kiểm tra; đơn giữ snapshot tại thời điểm đặt.'],
          ['Phân quyền','Token và role được kiểm tra ở API; khách chỉ truy cập đơn và thông báo thuộc quyền mình.'],
          ['Nghiệm thu','Chỉ đánh dấu đạt khi có kết quả kiểm tra gắn với phiên bản và môi trường cụ thể.'],
        ],[3.3,13.7])
    if index==3:
        heading('5 1 Số liệu đã lưu trong repository',2,newpage=True)
        para(ps[3])
        table(['Nội dung','Số liệu','Cách diễn giải'],[
          ['Catalog','89 active','Snapshot artifact local; chưa đếm lại DB live.'],
          ['Thông số','3.277 tổng\n1.012 cấu trúc\n2.265 text','Giữ text để bảo toàn ý nghĩa và nguồn.'],
          ['Benchmark v1.1','100 query\n93 tính metrics\n7 query âm','Đánh giá theo bộ nhãn và cấu hình của lần chạy.'],
          ['BGE Top 5 offline','nDCG@5\n0,941883 → 0,948655','Chất lượng thứ hạng trên benchmark; không phải xác suất đúng.'],
          ['Explanation','160 lượt HTTP\n326 accepted LLM','326 là số lời giải thích, không phải số request.'],
        ],[3.5,4.1,9.4])
        para('Các số liệu phản ánh những lần đo đã lưu. Không gộp số test giữa các lần chạy hoặc suy ra độ chính xác tuyệt đối cho phiên bản chưa được đánh giá lại.',italic=True,size=10.5)
    secnum=index+2
    for sub,number in enumerate(numbers,start=2):addfigure(number,secnum,sub)

newsection(False)
heading('6 Kết luận và bước tiếp theo')
conclusion=between('### Kết luận và bước tiếp theo','Trân trọng,')
for p in paragraphs(conclusion):para(p)
heading('6 1 Công việc trước buổi trình diễn',2)
for value in [
    'Chốt người phụ trách, phiên bản code, dữ liệu và cấu hình dùng làm mốc.',
    'FE kiểm tra client và admin; BE kiểm tra auth, giá, đơn, thanh toán và quyền.',
    'AI kiểm tra coverage embedding, cấu hình model, warmup và search smoke.',
    'BA cập nhật ma trận yêu cầu và ghi nhận đạt hoặc chưa đạt theo bằng chứng.',
    'Kiểm tra môi trường admin vì staging mặc định chặn thao tác ghi; chuẩn bị phần minh họa phù hợp.',
    'Tập trình bày theo FE → BE → BA → AI với bấm giờ và phương án dự phòng.',
]:doc.add_paragraph(value,style='List Bullet')
heading('6 2 Giới hạn cần quản lý',2)
para('AI Engineer kiêm dữ liệu có phạm vi lớn và cần hỗ trợ từ BE về hạ tầng, BA về tài liệu và tiêu chí. Model local có thể ảnh hưởng độ trễ. Đề xuất nhiều thành phần chưa xác nhận tương thích toàn hệ thống. Thanh toán hiện là VNPay sandbox; trạng thái deployment và chất lượng runtime cần được kiểm tra lại trước khi nghiệm thu.')

heading('Phụ lục A Kịch bản trình bày',newpage=True)
para('Thứ tự trình bày là FE → BE → BA → AI Engineer. Ngân sách 20 phút gồm minh họa ngắn, chưa gồm phản biện. Nhóm cần tập và điều chỉnh tốc độ nói theo thời lượng thực tế.')
timing=markdown_table(between('### 1. Thời lượng đề xuất','### 2. Lời trình bày của FE'))
table(timing[0],timing[1:],[2.7,3.4,10.9])
for i,(role,start,end) in enumerate([
('Frontend Developer','### 2. Lời trình bày của FE','### 3. Lời trình bày của BE'),
('Backend Developer','### 3. Lời trình bày của BE','### 4. Lời trình bày của BA'),
('Business Analyst','### 4. Lời trình bày của BA','### 5. Lời trình bày của AI Engineer'),
('AI Engineer và xử lý dữ liệu','### 5. Lời trình bày của AI Engineer','## D.'),
]):
    heading(role,2,newpage=i>0)
    content=between(start,end).split('\n',1)[1]
    chunks=[];current=[]
    for line in content.splitlines():
        if line.startswith('>'):
            val=line[1:].strip()
            if val:current.append(val)
            elif current:chunks.append(' '.join(current));current=[]
    if current:chunks.append(' '.join(current))
    for chunk in chunks:para(chunk)
    demo=re.search(r'\*\*Minh họa:\*\*\s*(.*)',content)
    if demo:para('Minh họa đề xuất  '+clean(demo.group(1)),italic=True,size=10.5)

heading('Phụ lục B Nguồn đối chiếu',newpage=True)
para('Các nguồn dưới đây nằm trong repository REECO AI Search. Tài liệu mô tả hiện trạng có thể có các đoạn lịch sử; khi khác nhau cần đối chiếu mã nguồn và report tại phiên bản đang xem.')
source_items=[
 ('Giao diện và điều hướng','apps/client/src/App.js; apps/admin/src/App.js'),
 ('Xác thực và phân quyền','apps/api/src/modules/auth/; apps/api/src/auth/'),
 ('Giá đơn và VNPay','apps/api/src/modules/admin/; apps/api/src/modules/orders/'),
 ('Database','supabase/migrations/'),
 ('Pipeline dữ liệu','scripts/product-master/; scripts/spec-normalization/; scripts/semantic-enrichment/; scripts/search-documents/; scripts/create-embeddings/'),
 ('Tìm kiếm và giải thích','apps/api/src/modules/search/; apps/api/src/ai/; services/reranker/'),
 ('Dữ liệu và benchmark','data/reports/product-master-report.json; data/reports/spec-normalization-report.json; docs/search-quality-baseline-v1.1.md; docs/reranker-bge-evaluation-v1.md'),
 ('Explanation HTTP','output/explanation-cases-complete/report.md'),
]
table(['Nội dung','Nguồn'],[list(x) for x in source_items],[4.4,12.6])
heading('Bản gốc sơ đồ có thể chỉnh sửa',2)
para('Các hình được chèn trực tiếp trong báo cáo. Bản gốc để chỉnh từng ô và đường nối nằm tại docs/diagrams/; file reeco-tat-ca-so-do.drawio có đủ 12 trang. Các file riêng được tổ chức theo FE, BE, BA và AI; ERD có hai trang catalog và commerce.')
para('Các thay đổi nội dung báo cáo thực hiện trực tiếp trong Word. Khi thay sơ đồ, chỉnh file draw.io tương ứng, xuất hình và thay ảnh trong tài liệu; giữ nguyên caption hoặc cập nhật danh mục hình cho phù hợp.',size=10.5)

# Word updates the TOC and page references during the native render pass.
settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
for p in doc.paragraphs:
    if p.style.name=='Heading 2':
        value=re.sub(r'^\d+(?:\s+\d+)+\s+','',p.text)
        if value!=p.text:p.runs[0].text=value
doc.save(OUTPUT)
with zipfile.ZipFile(OUTPUT) as z:
    assert len([n for n in z.namelist() if n.startswith('word/media/')])==12
    xml=z.read('word/document.xml').decode('utf-8')
    assert 'Sơ đồ minh họa:' not in xml and 'Sơ đồ chính:' not in xml
    assert '\ufffd' not in xml
print(json.dumps({'output':str(OUTPUT),'embedded_diagrams':12,'sections':len(doc.sections),'paragraphs':len(doc.paragraphs)},ensure_ascii=True))
